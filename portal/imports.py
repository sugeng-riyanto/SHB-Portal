"""Bounded XLSX preview. No database mutations occur here."""
from io import BytesIO
from datetime import date, datetime
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from zipfile import BadZipFile, ZipFile
from openpyxl import load_workbook, Workbook

HEADERS = ('nis', 'kelas', 'nama_murid', 'email_wali', 'jenis_tagihan', 'nominal', 'denda', 'deskripsi', 'tahun_ajaran', 'periode', 'jatuh_tempo', 'referensi_tagihan', 'referensi_rekening', 'nomor_va')
MAX_BYTES = 2 * 1024 * 1024
MAX_EXPANDED = 10 * 1024 * 1024
MAX_ROWS = 2000

class ImportErrorDetail(ValueError):
    pass

def template_bytes():
    book = Workbook()
    sheet = book.active
    sheet.title = 'Tagihan'
    sheet.append(HEADERS)
    for column in ('A', 'B', 'D', 'L', 'M', 'N'):
        sheet.column_dimensions[column].width = 24
    output = BytesIO()
    book.save(output)
    return output.getvalue()

def preview_xlsx(data):
    if not data or len(data) > MAX_BYTES:
        raise ImportErrorDetail('Ukuran file maksimal 2 MB.')
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(x.file_size for x in entries) > MAX_EXPANDED:
                raise ImportErrorDetail('Isi file terlalu besar.')
        book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            sheet = book.active
            if sheet.max_column > 50 or sheet.max_row > MAX_ROWS + 1:
                raise ImportErrorDetail('File melebihi batas 2.000 baris atau 50 kolom.')
            rows = sheet.iter_rows()
            header = tuple(str(c.value or '').strip().lower() for c in next(rows))
            if header[:len(HEADERS)] != HEADERS or any(header[len(HEADERS):]):
                raise ImportErrorDetail('Kolom tidak sesuai template.')
            results = []
            seen_references = set()
            for row_number, cells in enumerate(rows, start=2):
                extra_cells = cells[len(HEADERS):]
                cells = cells[:len(HEADERS)]
                if all(c.value is None for c in cells) and not any(c.value not in (None, '') for c in extra_cells):
                    continue
                errors = []
                if any(c.value not in (None, '') for c in extra_cells):
                    errors.append('Data di luar kolom template tidak diperbolehkan.')
                if any(c.data_type == 'f' for c in cells):
                    errors.append('Formula tidak diperbolehkan.')
                values = [c.value for c in cells]
                for i in (0, 1, 2, 3, 4, 8, 9, 10, 11):
                    if values[i] is None or not str(values[i]).strip():
                        errors.append(f'{HEADERS[i]} wajib diisi.')
                for i in (0, 11, 12, 13):
                    if values[i] is not None and not isinstance(values[i], str):
                        errors.append(f'{HEADERS[i]} harus berupa teks untuk menjaga nol di depan.')
                for i in (5, 6):
                    v = values[i]
                    if i == 6 and v is None:
                        continue
                    if isinstance(v, bool) or not isinstance(v, int) or v < 0:
                        errors.append(f'{HEADERS[i]} harus bilangan bulat rupiah nonnegatif.')
                email = values[3]
                if email:
                    try:
                        validate_email(str(email).strip())
                    except ValidationError:
                        errors.append('email_wali tidak valid.')
                due = values[10]
                if due and not isinstance(due, (date, datetime)):
                    try:
                        date.fromisoformat(str(due).strip())
                    except ValueError:
                        errors.append('jatuh_tempo harus berformat YYYY-MM-DD.')
                reference = values[11]
                if isinstance(reference, str) and reference.strip():
                    reference = reference.strip()
                    if reference in seen_references:
                        errors.append('referensi_tagihan duplikat dalam file.')
                    seen_references.add(reference)
                for i, limit in ((0,32),(1,20),(2,200),(3,254),(4,100),(7,200),(8,20),(9,100),(11,100),(12,64),(13,64)):
                    if values[i] is not None and len(str(values[i]).strip()) > limit:
                        errors.append(f'{HEADERS[i]} terlalu panjang.')
                if values[13]:
                    errors.append('Nomor VA belum dapat diverifikasi terhadap rekening sekolah; impor ditolak.')
                for i in (5,6):
                    if isinstance(values[i], int) and values[i] > 9223372036854775807:
                        errors.append(f'{HEADERS[i]} melebihi batas.')
                if isinstance(due, (date, datetime)): due = due.date().isoformat() if isinstance(due, datetime) else due.isoformat()
                results.append({'row':row_number,'nis':str(values[0] or ''),'errors':errors,
                    'classroom':str(values[1] or '').strip(), 'name':str(values[2] or '').strip(),
                    'email':str(values[3] or '').strip(), 'year':str(values[8] or '').strip(),
                    'reference':str(values[11] or '').strip(), 'bill_type':str(values[4] or '').strip(),
                    'amount':values[5] if isinstance(values[5],int) and not isinstance(values[5],bool) else 0,
                    'fine':values[6] if isinstance(values[6],int) and not isinstance(values[6],bool) else 0,
                    'description':str(values[7] or '').strip(), 'period':str(values[9] or '').strip(),
                    'due_date':str(due or '').strip(), 'bank_reference':str(values[12] or '').strip(), 'va':str(values[13] or '').strip()})
            return results
        finally:
            book.close()
    except ImportErrorDetail:
        raise
    except (BadZipFile, ValueError, KeyError, IndexError, StopIteration, OSError):
        raise ImportErrorDetail('File XLSX tidak valid.') from None

def match_rows(rows, school, lock_banks=False):
    """Check references within one authorized school; never create or link accounts."""
    from .models import BankAccount, Enrollment, Invoice, Membership
    for row in sorted(rows, key=lambda r: r.get('bank_reference','')):
        row.pop('enrollment_id', None)
        previous_bank = row.pop('bank_snapshot', None)
        if row.get('bank_reference'):
            accounts = BankAccount.objects.filter(school=school,active=True,account_number=row['bank_reference']).order_by('pk')
            if lock_banks: accounts = accounts.select_for_update()
            matches = list(accounts[:2])
            if len(matches) != 1:
                row['errors'].append('Rekening aktif tidak ditemukan secara unik pada sekolah.')
            else:
                account = matches[0]
                snapshot={'id':account.pk,'bank':account.bank_name,'name':account.account_name,'number':account.account_number}
                if lock_banks and previous_bank != snapshot:
                    row['errors'].append('Rekening berubah setelah preview; unggah ulang.')
                row['bank_snapshot'] = snapshot
        enrollment = Enrollment.objects.filter(student__school=school, student__active=True,
            student__nis=row['nis'].strip(), classroom__year__school=school,
            classroom__year__label=row['year'], classroom__name=row['classroom']).select_related('student')
        if enrollment.count() != 1:
            row['errors'].append('Murid, kelas dan tahun ajaran tidak cocok secara unik.')
        else:
            item = enrollment.first()
            student = item.student
            row['enrollment_id'] = item.pk
            if student.name.strip().casefold() != row['name'].casefold():
                row['errors'].append('Nama murid tidak cocok dengan data sekolah.')
            if not student.guardians.filter(is_active=True, email__iexact=row['email'],
                    membership__school=school, membership__active=True,
                    membership__role=Membership.Role.GUARDIAN).exists():
                row['errors'].append('Email belum cocok dengan akun wali aktif yang terhubung.')
        if row['reference'] and Invoice.objects.filter(reference=row['reference']).exists():
            row['errors'].append('Referensi tagihan sudah digunakan; perlu ditinjau.')
    return rows
