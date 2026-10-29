from io import BytesIO
from django import forms
from django.contrib.auth.decorators import login_required
from django.db.models import Q
from django.http import HttpResponse, Http404
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from .models import Membership, School, Student


def admin_school_ids(user):
    if not user.is_active:
        return School.objects.none().values_list('id', flat=True)
    if user.is_superuser:
        return School.objects.values_list('id', flat=True)
    return Membership.objects.filter(user=user, active=True, role=Membership.Role.ADMIN).values_list('school_id', flat=True)

class StudentForm(forms.ModelForm):
    class Meta:
        model = Student
        fields = ('school', 'nis', 'name', 'active')
        labels = {'school':'Sekolah','nis':'NIS','name':'Nama murid','active':'Aktif'}
    def __init__(self, *args, user, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields['school'].queryset = School.objects.filter(id__in=admin_school_ids(user))
    def clean_nis(self):
        return self.cleaned_data['nis'].strip()

@login_required
def student_list(request):
    schools = admin_school_ids(request.user)
    if not schools.exists():
        raise Http404
    query = request.GET.get('q','').strip()[:100]
    students = Student.objects.filter(school_id__in=schools).select_related('school')
    if query:
        students = students.filter(Q(name__icontains=query) | Q(nis__icontains=query))
    return render(request,'portal/student_list.html',{'students':students.order_by('name')[:200], 'query':query})

@login_required
def student_edit(request, student_id=None):
    schools = admin_school_ids(request.user)
    if not schools.exists():
        raise Http404
    instance = get_object_or_404(Student, pk=student_id,school_id__in=schools) if student_id else None
    form = StudentForm(request.POST or None, instance=instance, user=request.user)
    if request.method == 'POST' and form.is_valid():
        if instance and form.cleaned_data['school'].id != instance.school_id:
            form.add_error('school','Pemindahan sekolah memerlukan alur khusus.')
        else:
            form.save()
            return redirect('student_list')
    return render(request,'portal/student_form.html',{'form':form,'student':instance})

@login_required
@require_POST
def student_deactivate(request, student_id):
    student = get_object_or_404(Student,pk=student_id,school_id__in=admin_school_ids(request.user))
    student.active = False
    student.save(update_fields=['active'])
    return redirect('student_list')

def safe_cell(value):
    value = str(value or '')
    return "'" + value if value.lstrip().startswith(('=', '+', '-', '@')) else value

@login_required
def student_export(request):
    schools=admin_school_ids(request.user)
    if not schools.exists():
        raise Http404
    query=request.GET.get('q','').strip()[:100]
    rows=Student.objects.filter(school_id__in=schools).select_related('school')
    if query:
        rows=rows.filter(Q(name__icontains=query)|Q(nis__icontains=query))
    book=Workbook();sheet=book.active;sheet.title='Murid';sheet.append(['Sekolah','NIS','Nama','Status'])
    for student in rows.order_by('name').iterator():
        sheet.append([safe_cell(student.school.name),safe_cell(student.nis),safe_cell(student.name),'Aktif' if student.active else 'Nonaktif'])
    out=BytesIO();book.save(out)
    response=HttpResponse(out.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition']='attachment; filename="murid.xlsx"'
    response['Cache-Control']='private, no-store'
    return response

# Import only the fields currently supported by the student model. Class and
# guardian relationships require their own verified workflow.
from datetime import timedelta
from zipfile import BadZipFile, ZipFile
from django.db import IntegrityError, transaction
from django.http import HttpResponseBadRequest
from django.utils import timezone
from django.views.decorators.http import require_POST
from openpyxl import load_workbook
from .imports import MAX_BYTES, MAX_EXPANDED, MAX_ROWS

STUDENT_HEADERS=('nis','nama_murid')

@login_required
def student_template(request):
    if not admin_school_ids(request.user).exists(): raise Http404
    book=Workbook();sheet=book.active;sheet.title='Murid';sheet.append(STUDENT_HEADERS)
    sheet.column_dimensions['A'].width=24;sheet.column_dimensions['B'].width=36
    sheet.column_dimensions['A'].number_format='@'
    output=BytesIO();book.save(output)
    response=HttpResponse(output.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition']='attachment; filename="template-murid.xlsx"'
    return response

def read_student_xlsx(data):
    if not data or len(data)>MAX_BYTES: raise ValueError('Ukuran file maksimal 2 MB.')
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries=archive.infolist()
            if len(entries)>1000 or sum(entry.file_size for entry in entries)>MAX_EXPANDED:
                raise ValueError('Isi file terlalu besar.')
        book=load_workbook(BytesIO(data),read_only=True,data_only=False,keep_links=False)
        try:
            sheet=book.active
            if sheet.max_row>MAX_ROWS+1 or sheet.max_column>10: raise ValueError('Maksimal 2.000 baris dan 10 kolom.')
            rows=sheet.iter_rows();header=next(rows)
            if tuple(str(c.value or '').strip().lower() for c in header[:2]) != STUDENT_HEADERS:
                raise ValueError('Kolom harus nis dan nama_murid sesuai template.')
            result=[];seen=set()
            for number,cells in enumerate(rows,start=2):
                if all(c.value is None for c in cells): continue
                values=[c.value for c in cells[:2]]
                if len(values)<2: values += [None]*(2-len(values))
                errors=[]
                if any(c.data_type=='f' for c in cells): errors.append('Formula tidak diperbolehkan.')
                nis=values[0].strip() if isinstance(values[0],str) else ''
                name=values[1].strip() if isinstance(values[1],str) else ''
                if not nis or len(nis)>32: errors.append('NIS wajib berupa teks maksimal 32 karakter.')
                if not name or len(name)>200: errors.append('Nama wajib berupa teks maksimal 200 karakter.')
                if nis and nis in seen: errors.append('NIS duplikat di file.')
                seen.add(nis)
                result.append({'row':number,'nis':nis,'name':name,'errors':errors})
            if not result: raise ValueError('File tidak berisi data murid.')
            return result
        finally: book.close()
    except (BadZipFile, KeyError, IndexError, StopIteration, OSError, TypeError):
        raise ValueError('File XLSX tidak valid.') from None

@login_required
def student_import(request):
    schools=School.objects.filter(id__in=admin_school_ids(request.user))
    if not schools.exists(): raise Http404
    context={'schools':schools}
    if request.method=='POST':
        request.session.pop('student_import_preview',None)
        school=get_object_or_404(schools,pk=request.POST.get('school'))
        uploaded=request.FILES.get('file')
        if not uploaded or uploaded.size>MAX_BYTES or not uploaded.name.lower().endswith('.xlsx'):
            context['error']='Pilih file XLSX maksimal 2 MB.'
        else:
            try:
                rows=read_student_xlsx(uploaded.read(MAX_BYTES+1))
                existing=set(Student.objects.filter(school=school,nis__in=[r['nis'] for r in rows]).values_list('nis',flat=True))
                for row in rows:
                    if row['nis'] in existing: row['errors'].append('NIS sudah ada di sekolah ini.')
                context.update(rows=rows,school=school,can_commit=all(not row['errors'] for row in rows))
                if context['can_commit']:
                    import secrets
                    token=secrets.token_urlsafe(24)
                    request.session['student_import_preview']={'school':school.pk,'rows':rows,'token':token,'created':timezone.now().isoformat()}
                    context['commit_token']=token
            except ValueError as exc: context['error']=str(exc)
    return render(request,'portal/student_import.html',context)

@login_required
@require_POST
def student_import_commit(request):
    from datetime import datetime
    import secrets
    batch=request.session.get('student_import_preview')
    if not batch or not secrets.compare_digest(str(request.POST.get('token','')),batch['token']):
        return HttpResponseBadRequest('Preview tidak tersedia atau sudah digunakan.')
    try: created=datetime.fromisoformat(batch['created'])
    except (ValueError,TypeError): return HttpResponseBadRequest('Preview tidak valid.')
    if timezone.now()-created > timedelta(minutes=20):
        request.session.pop('student_import_preview',None)
        return HttpResponseBadRequest('Preview kedaluwarsa. Unggah ulang.')
    school=get_object_or_404(School,pk=batch['school'],pk__in=admin_school_ids(request.user))
    try:
        with transaction.atomic():
            School.objects.select_for_update().get(pk=school.pk)
            if Student.objects.filter(school=school,nis__in=[r['nis'] for r in batch['rows']]).exists():
                return HttpResponseBadRequest('NIS sudah ada. Unggah ulang untuk memeriksa data terbaru.')
            Student.objects.bulk_create([Student(school=school,nis=r['nis'],name=r['name']) for r in batch['rows']])
    except IntegrityError:
        return HttpResponseBadRequest('Data berubah atau duplikat. Unggah ulang untuk memeriksa kembali.')
    request.session.pop('student_import_preview',None)
    return redirect('student_list')
