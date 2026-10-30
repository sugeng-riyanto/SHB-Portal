from datetime import date
from django.contrib.auth import get_user_model
from django.test import TestCase, TransactionTestCase
from django.urls import reverse
from .models import School, Membership, AcademicYear, Classroom, Student, Enrollment, Invoice, TransferProof, Payment
import secrets as _secrets
TEST_PASSWORD_1 = 'pw-' + _secrets.token_urlsafe(24)
TEST_PASSWORD_2 = 'pw-' + _secrets.token_urlsafe(24)
TEST_PASSWORD_3 = 'pw-' + _secrets.token_urlsafe(24)
TEST_PASSWORD_4 = 'pw-' + _secrets.token_urlsafe(24)
TEST_PASSWORD_5 = 'pw-' + _secrets.token_urlsafe(24)
TEST_SECRET_6 = 'sk-' + _secrets.token_urlsafe(24)
TEST_PASSWORD_7 = 'pw-' + _secrets.token_urlsafe(24)
TEST_PASSWORD_8 = 'pw-' + _secrets.token_urlsafe(24)
TEST_PASSWORD_9 = 'pw-' + _secrets.token_urlsafe(24)
TEST_PASSWORD_10 = 'pw-' + _secrets.token_urlsafe(24)
TEST_PASSWORD_11 = 'pw-' + _secrets.token_urlsafe(24)
TEST_PASSWORD_12 = 'pw-' + _secrets.token_urlsafe(24)
TEST_PASSWORD_13 = 'pw-' + _secrets.token_urlsafe(24)


class DashboardAccessTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.guardian = User.objects.create_user(username='guardian',password=TEST_PASSWORD_3)
        self.teacher = User.objects.create_user(username='teacher',password=TEST_PASSWORD_3)
        self.other = User.objects.create_user(username='other',password=TEST_PASSWORD_3)
        self.school = School.objects.create(name='SHB')
        self.other_school = School.objects.create(name='Other')
        year = AcademicYear.objects.create(school=self.school,label='2026/27')
        self.classroom = Classroom.objects.create(year=year,name='8A')
        self.classroom.teachers.add(self.teacher)
        self.student = Student.objects.create(school=self.school,nis='001',name='ANAK SATU')
        self.student.guardians.add(self.guardian)
        self.guardian.email='a@example.test'; self.guardian.save(update_fields=['email'])
        enrollment = Enrollment.objects.create(student=self.student,classroom=self.classroom)
        self.invoice=Invoice.objects.create(enrollment=enrollment,reference='INV-1',description='SPP',amount=500000,due_date=date(2026,11,1))
        Membership.objects.create(user=self.guardian,school=self.school,role='guardian')
        Membership.objects.create(user=self.teacher,school=self.school,role='teacher')
        Membership.objects.create(user=self.other,school=self.other_school,role='admin')
    def test_login_required(self):
        self.assertRedirects(self.client.get('/dashboard/'),'/accounts/login/?next=/dashboard/')
    def test_linked_guardian_sees_invoice(self):
        self.client.force_login(self.guardian)
        self.assertContains(self.client.get('/dashboard/'),'ANAK SATU')
    def test_assigned_teacher_sees_invoice(self):
        self.client.force_login(self.teacher)
        self.assertContains(self.client.get('/dashboard/'),'ANAK SATU')
    def test_other_school_admin_cannot_see_invoice(self):
        self.client.force_login(self.other)
        self.assertNotContains(self.client.get('/dashboard/'),'ANAK SATU')
    def test_revoked_membership_denies_access(self):
        Membership.objects.filter(user=self.guardian).update(active=False)
        self.client.force_login(self.guardian)
        self.assertNotContains(self.client.get('/dashboard/'),'ANAK SATU')
    def test_logout_requires_post(self):
        self.client.force_login(self.guardian)
        self.assertEqual(self.client.get(reverse('logout')).status_code,405)

    def test_principal_and_foundation_see_only_aggregate(self):
        User = get_user_model()
        for role in ('principal', 'foundation'):
            user = User.objects.create_user(username=role,password=TEST_PASSWORD_3)
            Membership.objects.create(user=user,school=self.school,role=role)
            self.client.force_login(user)
            response = self.client.get('/dashboard/')
            self.assertNotContains(response, 'ANAK SATU')
            self.assertContains(response, '500000')
    def test_wrong_school_principal_has_no_aggregate(self):
        user = get_user_model().objects.create_user(username='otherprincipal',password=TEST_PASSWORD_3)
        Membership.objects.create(user=user,school=self.other_school,role='principal')
        self.client.force_login(user)
        self.assertNotContains(self.client.get('/dashboard/'), '500000')
    def test_other_guardian_and_unassigned_teacher_cannot_see(self):
        User = get_user_model()
        for role in ('guardian', 'teacher'):
            user = User.objects.create_user(username='unassigned-'+role,password=TEST_PASSWORD_3)
            Membership.objects.create(user=user,school=self.school,role=role)
            self.client.force_login(user)
            self.assertNotContains(self.client.get('/dashboard/'), 'ANAK SATU')
    def test_cross_school_invalid_enrollment_is_not_visible(self):
        year = AcademicYear.objects.create(school=self.other_school,label='2026/27')
        other_class = Classroom.objects.create(year=year,name='9A')
        other_class.teachers.add(self.teacher)
        bad = Enrollment.objects.create(student=self.student,classroom=other_class)
        Invoice.objects.create(enrollment=bad,reference='BAD',description='CROSS SCHOOL LEAK',amount=999,due_date=date(2026,11,1))
        self.client.force_login(self.guardian)
        self.assertNotContains(self.client.get('/dashboard/'),'CROSS SCHOOL LEAK')
        self.client.force_login(self.teacher)
        self.assertNotContains(self.client.get('/dashboard/'),'CROSS SCHOOL LEAK')
    def test_admin_sees_only_own_school(self):
        user = get_user_model().objects.create_user(username='adminshb',password=TEST_PASSWORD_3)
        Membership.objects.create(user=user,school=self.school,role='admin')
        self.client.force_login(user)
        self.assertContains(self.client.get('/dashboard/'),'ANAK SATU')

    def test_local_static_assets_and_viewport(self):
        response = self.client.get('/')
        self.assertContains(response, 'name="viewport"')
        self.assertContains(response, '/static/portal/app.css')
        self.assertContains(response, '/static/portal/app.js')
        self.assertNotContains(response, 'cdn.')

    def test_xlsx_import_requires_admin(self):
        self.client.force_login(self.guardian)
        self.assertEqual(self.client.get('/imports/').status_code,403)
        self.assertEqual(self.client.get('/imports/template.xlsx').status_code,403)
    def test_admin_download_and_preview(self):
        from io import BytesIO
        from django.core.files.uploadedfile import SimpleUploadedFile
        from openpyxl import load_workbook
        user = get_user_model().objects.create_user(username='admin_import',password=TEST_PASSWORD_3)
        Membership.objects.create(user=user,school=self.school,role='admin')
        self.client.force_login(user)
        template = self.client.get('/imports/template.xlsx')
        self.assertEqual(template.status_code,200)
        book = load_workbook(BytesIO(template.content))
        book.active.append(['001','8A','ANAK SATU','a@example.test','SPP',500000,0,'','2026/27','Oktober','2026-11-01','ref-001','',''])
        stream=BytesIO();book.save(stream)
        upload=SimpleUploadedFile('data.xlsx',stream.getvalue())
        result=self.client.post('/imports/',{'file':upload,'school':self.school.pk})
        self.assertContains(result,'Pemeriksaan awal lolos')
        self.assertContains(result,'001')
        self.assertEqual(Invoice.objects.count(),1)
    def test_formula_rejected_in_preview(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from .imports import template_bytes,preview_xlsx
        book=load_workbook(BytesIO(template_bytes()))
        book.active.append(['001','8A','Nama','a@example.test','SPP','=1+1',0,'','2026/27','Oktober','2026-11-01','ref','',''])
        stream=BytesIO();book.save(stream)
        self.assertTrue(any('Formula' in e for e in preview_xlsx(stream.getvalue())[0]['errors']))

    def test_guardian_upload_and_cross_family_privacy(self):
        from io import BytesIO
        from PIL import Image
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .models import TransferProof, Payment
        image=BytesIO();Image.new('RGB',(4,4),'white').save(image,format='PNG')
        self.client.force_login(self.guardian)
        response=self.client.post(f'/invoices/{self.invoice.pk}/proof/',{'amount':'200000','transfer_date':'2026-10-27','image':SimpleUploadedFile('proof.png',image.getvalue())})
        self.assertEqual(response.status_code,302)
        proof=TransferProof.objects.get()
        self.assertEqual(Payment.objects.count(),0)
        self.assertEqual(self.client.get(f'/proofs/{proof.id}/image/').status_code,200)
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(f'/proofs/{proof.id}/image/').status_code,404)
    def test_admin_verify_once_and_cross_role_denied(self):
        from .models import TransferProof, Payment
        proof=TransferProof.objects.create(invoice=Invoice.objects.get(),guardian=self.guardian,reported_amount=100000,transfer_date=date(2026,10,27),image=b'x',image_type='png')
        self.client.force_login(self.teacher)
        self.assertEqual(self.client.post(f'/proofs/{proof.id}/verify/').status_code,404)
        admin=get_user_model().objects.create_user(username='proof-admin',password=TEST_PASSWORD_3)
        Membership.objects.create(user=admin,school=self.school,role='admin')
        self.client.force_login(admin)
        self.assertEqual(self.client.post(f'/proofs/{proof.id}/verify/').status_code,302)
        self.assertEqual(self.client.post(f'/proofs/{proof.id}/verify/').status_code,400)
        self.assertEqual(Payment.objects.count(),1)
    def test_invalid_photo_rejected(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from .models import TransferProof
        self.client.force_login(self.guardian)
        response=self.client.post(f'/invoices/{self.invoice.pk}/proof/',{'amount':'100','transfer_date':'2026-10-27','image':SimpleUploadedFile('fake.png',b'not an image')})
        self.assertContains(response,'Foto tidak valid')
        self.assertEqual(TransferProof.objects.count(),0)

    def test_superadmin_can_manage_school_and_view_data(self):
        user = get_user_model().objects.create_superuser(username='root-user',email='root@example.test',password=TEST_PASSWORD_3)
        self.client.force_login(user)
        self.assertContains(self.client.get('/dashboard/'),'ANAK SATU')
        self.assertEqual(self.client.get('/admin/portal/school/').status_code,200)
        self.assertEqual(self.client.get('/admin/portal/membership/').status_code,200)
    def test_staff_without_superuser_cannot_use_admin(self):
        staff = get_user_model().objects.create_user(username='staff-user',password=TEST_PASSWORD_3,is_staff=True)
        self.client.force_login(staff)
        self.assertEqual(self.client.get('/admin/portal/membership/').status_code,302)

    def test_integration_status_only_superuser_never_displays_secret(self):
        from unittest.mock import patch
        root = get_user_model().objects.create_superuser(username='super-settings',email='root@example.test',password=TEST_PASSWORD_3)
        with patch.dict('os.environ', {'SMTP_PASSWORD':'secret-value-for-test'}):
            self.client.force_login(root)
            response = self.client.get('/superadmin/integrations/')
            self.assertContains(response,'Midtrans')
            self.assertNotContains(response,'secret-value-for-test')
            self.client.force_login(self.guardian)
            self.assertEqual(self.client.get('/superadmin/integrations/').status_code,404)

    def test_student_crud_scoped_to_admin_school(self):
        User=get_user_model()
        admin=User.objects.create_user(username='student-admin',password=TEST_PASSWORD_3)
        Membership.objects.create(user=admin,school=self.school,role='admin')
        outsider=Student.objects.create(school=self.other_school,nis='009',name='ANAK LUAR')
        self.client.force_login(admin)
        self.assertContains(self.client.get('/students/'),'ANAK SATU')
        self.assertNotContains(self.client.get('/students/'),'ANAK LUAR')
        self.assertEqual(self.client.get(f'/students/{outsider.pk}/edit/').status_code,404)
        self.assertEqual(self.client.post('/students/new/',{'school':self.other_school.pk,'nis':'002','name':'NOT ALLOWED','active':'on'}).status_code,200)
        self.assertFalse(Student.objects.filter(nis='002').exists())
        self.assertEqual(self.client.post('/students/new/',{'school':self.school.pk,'nis':'0002','name':'Anak baru','active':'on'}).status_code,302)
        self.assertTrue(Student.objects.filter(nis='0002').exists())
        self.assertEqual(self.client.post(f'/students/{self.student.pk}/deactivate/').status_code,302)
        self.student.refresh_from_db();self.assertFalse(self.student.active)
    def test_guardian_cannot_manage_students(self):
        self.client.force_login(self.guardian)
        self.assertEqual(self.client.get('/students/').status_code,404)
        self.assertEqual(self.client.get('/students/export.xlsx').status_code,404)
        self.assertEqual(self.client.post(f'/students/{self.student.pk}/deactivate/').status_code,404)
    def test_export_escapes_formulas_and_preserves_scope(self):
        from io import BytesIO
        from openpyxl import load_workbook
        admin=get_user_model().objects.create_user(username='export-admin',password=TEST_PASSWORD_3)
        Membership.objects.create(user=admin,school=self.school,role='admin')
        Student.objects.create(school=self.school,nis='00009',name='=1+1')
        Student.objects.create(school=self.other_school,nis='00010',name='PRIVATE')
        self.client.force_login(admin)
        response=self.client.get('/students/export.xlsx')
        self.assertEqual(response.status_code,200)
        rows=list(load_workbook(BytesIO(response.content)).active.values)
        self.assertIn(('SHB','00009',"'=1+1",'Aktif'),rows)
        self.assertFalse(any('PRIVATE' in row for row in rows))

    def test_bank_crud_principal_scope_and_export(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from .models import BankAccount
        principal=get_user_model().objects.create_user(username='bank-principal',password=TEST_PASSWORD_3)
        Membership.objects.create(user=principal,school=self.school,role='principal')
        external=BankAccount.objects.create(school=self.other_school,bank_name='Other',account_name='Private',account_number='098')
        self.client.force_login(principal)
        response=self.client.post('/banks/new/',{'school':self.school.pk,'bank_name':'Bank','account_name':'SHB','account_number':'00123','active':'on'})
        self.assertEqual(response.status_code,302)
        bank=BankAccount.objects.get(account_number='00123')
        self.assertEqual(self.client.get(f'/banks/{external.pk}/edit/').status_code,404)
        self.assertNotContains(self.client.get('/banks/'),'Private')
        rows=list(load_workbook(BytesIO(self.client.get('/banks/export.xlsx').content)).active.values)
        self.assertTrue(any('00123' in row for row in rows))
        self.assertFalse(any('Private' in row for row in rows))
        self.assertEqual(self.client.post(f'/banks/{bank.pk}/deactivate/').status_code,302)
        bank.refresh_from_db();self.assertFalse(bank.active)
    def test_guardian_and_admin_cannot_manage_banks(self):
        self.client.force_login(self.guardian)
        self.assertEqual(self.client.get('/banks/').status_code,404)
        admin=get_user_model().objects.create_user(username='bank-admin',password=TEST_PASSWORD_3)
        Membership.objects.create(user=admin,school=self.school,role='admin')
        self.client.force_login(admin)
        self.assertEqual(self.client.get('/banks/').status_code,404)

    def test_sensitive_proof_response_no_store(self):
        from .models import TransferProof
        proof=TransferProof.objects.create(invoice=Invoice.objects.get(),guardian=self.guardian,reported_amount=1,transfer_date=date(2026,10,27),image=b'x',image_type='png')
        self.client.force_login(self.guardian)
        response=self.client.get(f'/proofs/{proof.pk}/image/')
        self.assertEqual(response['Cache-Control'],'private, no-store')
        self.assertEqual(response['X-Content-Type-Options'],'nosniff')

    def test_admin_email_and_guardian_inbox(self):
        from django.core import mail
        from django.test import override_settings
        from .models import Message
        admin=get_user_model().objects.create_user(username='mailadmin',password=TEST_PASSWORD_3)
        Membership.objects.create(user=admin,school=self.school,role='admin')
        self.guardian.email='parent@example.test';self.guardian.save()
        self.client.force_login(admin)
        with override_settings(SHB_EMAIL_ENABLED=True,EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend',DEFAULT_FROM_EMAIL='school@example.test'):
            response=self.client.post('/messages/compose/',{'school':self.school.pk,'recipient':self.guardian.pk,'subject':'Informasi','body':'Halo wali'})
            self.assertEqual(response.status_code,302)
            self.assertEqual(len(mail.outbox),1)
            self.assertEqual(mail.outbox[0].to,['parent@example.test'])
            self.assertEqual(Message.objects.get().email_status,'sent')
        self.client.force_login(self.guardian)
        self.assertContains(self.client.get('/messages/inbox/'),'Halo wali')
        self.client.force_login(self.other)
        self.assertNotContains(self.client.get('/messages/inbox/'),'Halo wali')
    def test_email_failure_and_wrong_school_recipient(self):
        from django.test import override_settings
        from .models import Message
        admin=get_user_model().objects.create_user(username='mailadmin2',password=TEST_PASSWORD_3)
        Membership.objects.create(user=admin,school=self.other_school,role='admin')
        self.guardian.email='parent@example.test';self.guardian.save()
        self.client.force_login(admin)
        self.client.post('/messages/compose/',{'school':self.other_school.pk,'recipient':self.guardian.pk,'subject':'private','body':'message'})
        self.assertEqual(Message.objects.count(),0)
        Membership.objects.create(user=admin,school=self.school,role='admin')
        with override_settings(SHB_EMAIL_ENABLED=False):
            self.client.post('/messages/compose/',{'school':self.school.pk,'recipient':self.guardian.pk,'subject':'private','body':'message'})
        self.assertEqual(Message.objects.get().email_status,'failed')
        self.client.force_login(self.teacher)
        self.assertEqual(self.client.get('/messages/compose/').status_code,404)

    def test_seed_demo_is_idempotent_and_cannot_run_on_production(self):
        from django.db import connection
        if connection.vendor != 'sqlite':
            self.skipTest('seed_demo diizinkan hanya pada SQLite development.')
        from io import StringIO
        from django.core.management import call_command
        from django.core.management.base import CommandError
        from django.test import override_settings
        out=StringIO();call_command('seed_demo',stdout=out)
        user=get_user_model().objects.get(username='ortu@shb.sch.id')
        self.assertTrue(user.check_password('demo123'))
        user.set_password('different-password');user.save()
        call_command('seed_demo',stdout=StringIO())
        user.refresh_from_db();self.assertTrue(user.check_password('demo123'))
        with override_settings(PRODUCTION=True):
            with self.assertRaises(CommandError): call_command('seed_demo',stdout=StringIO())

    def test_csrf_login_trusted_origin_and_untrusted_origin(self):
        from django.test import Client, override_settings
        user = get_user_model().objects.create_user(username='csrfuser',password=TEST_PASSWORD_3)
        client = Client(enforce_csrf_checks=True)
        with override_settings(CSRF_TRUSTED_ORIGINS=['https://preview.example.test']):
            client.get('/accounts/login/')
            token = client.cookies['csrftoken'].value
            response=client.post('/accounts/login/',{'username':user.username,'password':TEST_PASSWORD_3,'csrfmiddlewaretoken':token},HTTP_ORIGIN='https://untrusted.example.test')
            self.assertEqual(response.status_code,403)
            response=client.post('/accounts/login/',{'username':user.username,'password':TEST_PASSWORD_3,'csrfmiddlewaretoken':token},HTTP_ORIGIN='https://preview.example.test')
            self.assertEqual(response.status_code,302)

    def test_demo_role_logins_and_dashboards_with_csrf(self):
        from django.db import connection
        if connection.vendor != 'sqlite':
            self.skipTest('seed_demo diizinkan hanya pada SQLite development.')
        from io import StringIO
        from django.core.management import call_command
        from django.test import Client, override_settings
        call_command('seed_demo', stdout=StringIO())
        origin='https://pv-ac9ad268c614-51b1306d-b202-4d34-b04b-346824493b23.style.dev'
        with override_settings(CSRF_TRUSTED_ORIGINS=[origin]):
            for role in ('superadmin','admin','kepala','yayasan','walikelas','ortu'):
                with self.subTest(role=role):
                    client=Client(enforce_csrf_checks=True)
                    client.get('/accounts/login/')
                    response=client.post('/accounts/login/',{'username':role+'@shb.sch.id','password':'demo123','csrfmiddlewaretoken':client.cookies['csrftoken'].value},HTTP_ORIGIN=origin)
                    self.assertEqual(response.status_code,302)
                    self.assertEqual(client.get('/dashboard/').status_code,302)
                    self.assertEqual(client.get('/accounts/password/').status_code,200)

    def test_class_recipient_selection_deduplicates_and_isolates(self):
        from .recipients import class_recipients
        admin=get_user_model().objects.create_user(username='bulkadmin',password=TEST_PASSWORD_3)
        Membership.objects.create(user=admin,school=self.school,role='admin')
        self.guardian.email='parent@example.test';self.guardian.save()
        second=Student.objects.create(school=self.school,nis='002',name='SECOND CHILD')
        second.guardians.add(self.guardian)
        Enrollment.objects.create(student=second,classroom=self.classroom)
        recipients=list(class_recipients(admin,self.classroom.pk))
        self.assertEqual([u.pk for u in recipients],[self.guardian.pk])
        self.assertFalse(class_recipients(self.other,self.classroom.pk).exists())
        self.assertFalse(class_recipients(self.teacher,self.classroom.pk).exists())

    def test_first_login_requires_password_change_and_preserves_session(self):
        user=get_user_model().objects.create_user(username='first',password=TEST_PASSWORD_7)
        self.assertTrue(self.client.login(username='first',password=TEST_PASSWORD_7))
        # Actual login view provisions a first-login profile.
        self.client.logout()
        self.client.post('/accounts/login/',{'username':'first','password':TEST_PASSWORD_7})
        self.assertRedirects(self.client.get('/dashboard/'),'/accounts/password/')
        response=self.client.post('/accounts/password/',{'old_password':TEST_PASSWORD_7,'new_password1':'Fresh-password-654!','new_password2':'Fresh-password-654!'})
        self.assertEqual(response.status_code,302)
        self.assertEqual(self.client.get('/dashboard/').status_code,200)
        user.refresh_from_db();self.assertTrue(user.check_password('Fresh-password-654!'))
    def test_user_cannot_see_other_profile_photo(self):
        from .models import UserProfile
        UserProfile.objects.create(user=self.guardian,must_change_password=False,photo=b'secret',photo_type='png')
        self.client.force_login(self.other)
        self.assertEqual(self.client.get('/accounts/profile/photo/').status_code,404)
    def test_login_has_local_password_toggle(self):
        response=self.client.get('/accounts/login/')
        self.assertContains(response,'data-password-target="id_password"')
        self.assertContains(response,'csrfmiddlewaretoken')

    def test_school_news_only_published_and_current(self):
        from datetime import timedelta
        from django.utils import timezone
        from .models import SchoolUpdate
        SchoolUpdate.objects.create(title='Published news',body='School news',published=True,published_at=timezone.now()-timedelta(days=1))
        SchoolUpdate.objects.create(title='PRIVATE DRAFT',body='Draft',published=False,published_at=timezone.now())
        SchoolUpdate.objects.create(title='FUTURE NEWS',body='Future',published=True,published_at=timezone.now()+timedelta(days=1))
        for url in ('/','/kabar-sekolah/'):
            response=self.client.get(url)
            self.assertContains(response,'Published news')
            self.assertNotContains(response,'PRIVATE DRAFT')
            self.assertNotContains(response,'FUTURE NEWS')
    def test_social_links_require_https_and_open_safely(self):
        from unittest.mock import patch
        with patch.dict('os.environ',{'SHB_INSTAGRAM_URL':'https://instagram.com/example','SHB_YOUTUBE_URL':'javascript:alert(1)'}):
            response=self.client.get('/')
            self.assertContains(response,'https://instagram.com/example')
            self.assertContains(response,'rel="noopener noreferrer"')
            self.assertNotContains(response,'javascript:alert(1)')

    def test_home_redirects_signed_in_users_to_dashboard(self):
        self.client.force_login(self.guardian)
        self.assertRedirects(self.client.get('/'),'/dashboard/')
    def test_login_lands_on_dashboard_even_if_next_provided(self):
        user=get_user_model().objects.create_user(username='landing-user',password=TEST_PASSWORD_7)
        response=self.client.post('/accounts/login/?next=/kabar-sekolah/',{'username':user.username,'password':TEST_PASSWORD_7})
        self.assertEqual(response.status_code,302)
        self.assertEqual(response['Location'],'/dashboard/')

    def test_language_selector_default_and_choices(self):
        response=self.client.get('/')
        self.assertContains(response,'Stay connected with school')
        self.client.cookies['django_language']='id'
        self.assertContains(self.client.get('/'),'Terhubung dengan sekolah')
        self.client.cookies['django_language']='zh-hant'
        self.assertContains(self.client.get('/'),'與學校保持聯繫')

    def test_proof_owner_can_edit_and_withdraw_before_verification(self):
        from .models import TransferProof
        proof=TransferProof.objects.create(invoice=Invoice.objects.get(),guardian=self.guardian,reported_amount=100000,transfer_date=date(2026,10,27),image=b'x',image_type='png')
        self.client.force_login(self.guardian)
        self.assertEqual(self.client.post(f'/proofs/{proof.pk}/edit/',{'amount':'120000','transfer_date':'2026-10-28'}).status_code,302)
        proof.refresh_from_db();self.assertEqual(proof.reported_amount,120000)
        self.assertEqual(self.client.post(f'/proofs/{proof.pk}/withdraw/').status_code,302)
        proof.refresh_from_db();self.assertEqual(proof.status,'withdrawn')
        self.assertEqual(self.client.get(f'/proofs/{proof.pk}/edit/').status_code,404)
    def test_other_guardian_cannot_edit_proof_and_verified_is_immutable(self):
        from .models import TransferProof
        proof=TransferProof.objects.create(invoice=Invoice.objects.get(),guardian=self.guardian,reported_amount=100000,transfer_date=date(2026,10,27),image=b'x',image_type='png')
        user=get_user_model().objects.create_user(username='other-wali',password=TEST_PASSWORD_3)
        Membership.objects.create(user=user,school=self.school,role='guardian');self.student.guardians.add(user)
        self.client.force_login(user)
        self.assertEqual(self.client.get(f'/proofs/{proof.pk}/edit/').status_code,404)
        self.assertEqual(self.client.post(f'/proofs/{proof.pk}/withdraw/').status_code,404)
        proof.status='verified';proof.save(update_fields=['status'])
        self.client.force_login(self.guardian)
        self.assertEqual(self.client.get(f'/proofs/{proof.pk}/edit/').status_code,404)
        self.assertEqual(self.client.post(f'/proofs/{proof.pk}/withdraw/').status_code,400)

    def test_xlsx_invalid_email_date_and_duplicate_reference(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from .imports import template_bytes, preview_xlsx
        book=load_workbook(BytesIO(template_bytes()))
        book.active.append(['001','8A','Student','bad-email','SPP',100,0,'','2026/27','Oktober','tomorrow','ref-1','',''])
        book.active.append(['002','8A','Student2','parent@example.test','SPP',100,0,'','2026/27','Oktober','2026-11-01','ref-1','',''])
        stream=BytesIO();book.save(stream)
        rows=preview_xlsx(stream.getvalue())
        self.assertTrue(any('email_wali' in e for e in rows[0]['errors']))
        self.assertTrue(any('jatuh_tempo' in e for e in rows[0]['errors']))
        self.assertTrue(any('duplikat' in e for e in rows[1]['errors']))

    def test_superadmin_integrations_encrypt_and_do_not_expose_secret(self):
        from cryptography.fernet import Fernet
        from unittest.mock import patch
        from .models import IntegrationConfig
        root=get_user_model().objects.create_superuser(username='integration-root',email='root@example.test',password=TEST_PASSWORD_3)
        self.client.force_login(root)
        key=Fernet.generate_key().decode()
        with patch.dict('os.environ', {'SHB_INTEGRATION_ENCRYPTION_KEY':key}):
            response=self.client.post('/superadmin/integrations/smtp/',{'host':'smtp.gmail.com','port':'587','username':'school@example.test','from_email':'school@example.test','app_password':'test-secret-123','active':'on'})
            self.assertEqual(response.status_code,302)
            config=IntegrationConfig.objects.get(provider='smtp')
            self.assertNotIn(b'test-secret-123',bytes(config.encrypted_secrets))
            response=self.client.get('/superadmin/integrations/smtp/')
            self.assertNotContains(response,'test-secret-123')
            self.assertEqual(response['Cache-Control'],'private, no-store')
            self.client.post('/superadmin/integrations/smtp/',{'host':'smtp.gmail.com','port':'587','username':'school@example.test','from_email':'school@example.test','app_password':'','active':'on'})
            config.refresh_from_db()
            self.assertIn(b'test-secret-123',Fernet(key.encode()).decrypt(bytes(config.encrypted_secrets)))
            self.assertEqual(self.client.post('/superadmin/integrations/smtp/remove/').status_code,302)
            self.assertFalse(IntegrationConfig.objects.exists())
        self.client.force_login(self.guardian)
        self.assertEqual(self.client.get('/superadmin/integrations/smtp/').status_code,404)
    def test_integration_missing_encryption_key_does_not_save(self):
        from unittest.mock import patch
        from .models import IntegrationConfig
        root=get_user_model().objects.create_superuser(username='nokey',email='root@example.test',password=TEST_PASSWORD_3)
        self.client.force_login(root)
        with patch.dict('os.environ', {'SHB_INTEGRATION_ENCRYPTION_KEY':''}):
            response=self.client.post('/superadmin/integrations/smtp/',{'host':'smtp.gmail.com','port':'587','username':'school@example.test','from_email':'school@example.test','app_password':'secret'})
        self.assertContains(response,'belum disimpan')
        self.assertFalse(IntegrationConfig.objects.exists())

    def test_superadmin_branding_and_local_assets(self):
        root=get_user_model().objects.create_superuser(username='styled-admin',email='root@example.test',password=TEST_PASSWORD_3)
        self.client.force_login(root)
        response=self.client.get('/admin/')
        self.assertContains(response,'SHB Portal')
        self.assertContains(response,'/static/portal/admin-shb.css')
        self.assertContains(response,'/superadmin/integrations/')
        self.assertNotContains(response,'cdn.')

    def test_secret_reveal_endpoint_is_not_available(self):
        root=get_user_model().objects.create_superuser(username='reveal-root',email='root@example.test',password=TEST_PASSWORD_2)
        self.client.force_login(root)
        self.assertEqual(self.client.get('/superadmin/integrations/smtp/reveal/').status_code,404)
        self.assertEqual(self.client.post('/superadmin/integrations/smtp/reveal/',{'password':TEST_PASSWORD_2}).status_code,404)

    def test_proof_upload_form_contains_csrf_and_accepts_valid_token(self):
        from io import BytesIO
        from PIL import Image
        from django.test import Client
        from django.core.files.uploadedfile import SimpleUploadedFile
        client=Client(enforce_csrf_checks=True);client.force_login(self.guardian)
        response=client.get(f'/invoices/{self.invoice.pk}/proof/')
        self.assertContains(response,'csrfmiddlewaretoken')
        token=client.cookies['csrftoken'].value
        image=BytesIO();Image.new('RGB',(3,3),'white').save(image,format='PNG')
        response=client.post(f'/invoices/{self.invoice.pk}/proof/',{'csrfmiddlewaretoken':token,'amount':'10000','transfer_date':'2026-10-27','image':SimpleUploadedFile('proof.png',image.getvalue())})
        self.assertEqual(response.status_code,302)
    def test_bulk_approval_is_atomic_and_denies_other_school(self):
        from .models import TransferProof, Payment
        invoice=Invoice.objects.get()
        p1=TransferProof.objects.create(invoice=invoice,guardian=self.guardian,reported_amount=100000,transfer_date=date(2026,10,27),image=b'x',image_type='png')
        p2=TransferProof.objects.create(invoice=invoice,guardian=self.guardian,reported_amount=500000,transfer_date=date(2026,10,27),image=b'x',image_type='png')
        admin=get_user_model().objects.create_user(username='bulk-proof',password=TEST_PASSWORD_3)
        Membership.objects.create(user=admin,school=self.school,role='admin')
        self.client.force_login(self.other)
        self.assertEqual(self.client.post('/proofs/bulk-verify/',{'proof_ids':[p1.pk]}).status_code,400)
        self.client.force_login(admin)
        self.assertEqual(self.client.post('/proofs/bulk-verify/',{'proof_ids':[p1.pk,p2.pk]}).status_code,400)
        self.assertEqual(Payment.objects.count(),0)
        p2.reported_amount=300000;p2.save(update_fields=['reported_amount'])
        self.assertEqual(self.client.post('/proofs/bulk-verify/',{'proof_ids':[p1.pk,p2.pk]}).status_code,302)
        self.assertEqual(Payment.objects.count(),2)
        self.assertEqual(self.client.post('/proofs/bulk-verify/',{'proof_ids':[p1.pk]}).status_code,400)
    def test_proof_comments_private_between_owner_and_admin(self):
        from .models import TransferProof, ProofComment
        proof=TransferProof.objects.create(invoice=Invoice.objects.get(),guardian=self.guardian,reported_amount=100,transfer_date=date(2026,10,27),image=b'x',image_type='png')
        self.client.force_login(self.guardian)
        self.assertEqual(self.client.post(f'/proofs/{proof.pk}/comments/',{'body':'Mohon diperiksa'}).status_code,302)
        self.client.force_login(self.teacher)
        self.assertEqual(self.client.get(f'/proofs/{proof.pk}/comments/').status_code,404)
        admin=get_user_model().objects.create_user(username='comment-admin',password=TEST_PASSWORD_3)
        Membership.objects.create(user=admin,school=self.school,role='admin')
        self.client.force_login(admin)
        self.assertContains(self.client.get(f'/proofs/{proof.pk}/comments/'),'Mohon diperiksa')
        self.assertEqual(ProofComment.objects.count(),1)

    def test_responsive_navigation_has_accessible_toggle(self):
        response=self.client.get('/')
        self.assertContains(response,'data-menu-toggle')
        self.assertContains(response,'aria-controls="main-navigation"')
        self.assertContains(response,'id="main-navigation"')
        self.assertContains(response,'/static/portal/app.js')

    def test_demo_data_works_for_guardian_and_hides_upload_for_teacher(self):
        from django.db import connection
        # seed_demo refuses to run outside a development SQLite database by design.
        if connection.vendor != 'sqlite':
            self.skipTest('seed_demo diizinkan hanya pada SQLite development.')
        from io import StringIO
        from django.core.management import call_command
        from .models import Invoice
        call_command('seed_demo',stdout=StringIO())
        call_command('seed_demo_data',stdout=StringIO())
        call_command('seed_demo_data',stdout=StringIO())
        invoice=Invoice.objects.get(reference='SHB-DEMO-INV-001')
        guardian=get_user_model().objects.get(username='ortu@shb.sch.id')
        teacher=get_user_model().objects.get(username='walikelas@shb.sch.id')
        self.client.force_login(guardian)
        self.assertContains(self.client.get('/dashboard/'),f'/invoices/{invoice.pk}/proof/')
        self.assertEqual(self.client.get(f'/invoices/{invoice.pk}/proof/').status_code,200)
        self.client.force_login(teacher)
        self.assertNotContains(self.client.get('/dashboard/'),f'/invoices/{invoice.pk}/proof/')
        self.assertEqual(self.client.get(f'/invoices/{invoice.pk}/proof/').status_code,404)

    def test_payment_types_defaults_crud_and_scope(self):
        from .models import PaymentType,ActivityEvent
        admin=get_user_model().objects.create_user(username='types-admin',password=TEST_PASSWORD_3)
        Membership.objects.create(user=admin,school=self.school,role='admin')
        self.client.force_login(admin)
        self.assertEqual(self.client.post('/payment-types/defaults/',{'school':self.school.pk}).status_code,302)
        self.client.post('/payment-types/defaults/',{'school':self.school.pk})
        self.assertEqual(PaymentType.objects.filter(school=self.school).count(),5)
        self.assertEqual(set(PaymentType.objects.filter(school=self.school).values_list('name',flat=True)),{'Monthly','ECA','Annual Fee','Stationery','Late Fee'})
        self.assertEqual(self.client.post('/payment-types/new/',{'school':self.school.pk,'name':'Books','active':'on'}).status_code,302)
        item=PaymentType.objects.get(school=self.school,name='Books')
        self.assertEqual(self.client.post(f'/payment-types/{item.pk}/deactivate/').status_code,302)
        item.refresh_from_db();self.assertFalse(item.active)
        self.assertTrue(ActivityEvent.objects.filter(school=self.school,actor=admin).exists())
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(f'/payment-types/{item.pk}/edit/').status_code,404)
        self.assertNotContains(self.client.get('/activities/'),'types-admin')
        self.client.force_login(self.guardian)
        self.assertEqual(self.client.get('/payment-types/').status_code,404)
        self.assertEqual(self.client.get('/activities/').status_code,404)

    def test_dashboard_billing_timeline_respects_scope(self):
        self.client.force_login(self.guardian)
        self.assertContains(self.client.get('/dashboard/'),'Billing timeline')
        self.assertContains(self.client.get('/dashboard/'),'Nov 2026')
        self.client.force_login(self.other)
        self.assertNotContains(self.client.get('/dashboard/'),'Nov 2026')

    def test_invoice_export_never_exposes_other_school(self):
        from io import BytesIO
        from openpyxl import load_workbook
        self.client.force_login(self.guardian)
        response=self.client.get('/dashboard/export.xlsx')
        self.assertEqual(response.status_code,200)
        self.assertIn('no-store',response['Cache-Control'])
        rows=list(load_workbook(BytesIO(response.content)).active.values)
        self.assertTrue(any('ANAK SATU' in row for row in rows))
        self.client.force_login(self.other)
        rows=list(load_workbook(BytesIO(self.client.get('/dashboard/export.xlsx').content)).active.values)
        self.assertFalse(any('ANAK SATU' in row for row in rows))

    def test_superadmin_and_school_admin_dashboard_are_distinct(self):
        User=get_user_model()
        root=User.objects.create_superuser(username='dash-super',email='root@example.test',password=TEST_PASSWORD_3)
        admin=User.objects.create_user(username='dash-admin',password=TEST_PASSWORD_3)
        Membership.objects.create(user=admin,school=self.school,role='admin')
        self.client.force_login(root)
        response=self.client.get('/dashboard/')
        self.assertContains(response,'Super admin dashboard')
        self.assertContains(response,'All schools')
        self.assertContains(response,'/superadmin/integrations/')
        self.client.force_login(admin)
        response=self.client.get('/dashboard/')
        self.assertContains(response,'School admin dashboard')
        self.assertContains(response,'Assigned schools')
        self.assertNotContains(response,'/superadmin/integrations/')
        self.assertContains(response,'500000')
        self.assertEqual(self.client.get('/superadmin/integrations/').status_code,404)

    def test_receipt_only_when_invoice_fully_paid_and_scoped(self):
        from .models import TransferProof,PaymentReceipt
        invoice=Invoice.objects.get()
        proof=TransferProof.objects.create(invoice=invoice,guardian=self.guardian,reported_amount=200000,transfer_date=date(2026,10,27),image=b'x',image_type='png')
        admin=get_user_model().objects.create_user(username='receipt-admin',password=TEST_PASSWORD_3)
        Membership.objects.create(user=admin,school=self.school,role='admin')
        self.client.force_login(admin)
        self.assertEqual(self.client.post(f'/proofs/{proof.pk}/verify/').status_code,302)
        self.assertFalse(PaymentReceipt.objects.exists())
        proof2=TransferProof.objects.create(invoice=invoice,guardian=self.guardian,reported_amount=300000,transfer_date=date(2026,10,27),image=b'x',image_type='png')
        self.assertEqual(self.client.post(f'/proofs/{proof2.pk}/verify/').status_code,302)
        receipt=PaymentReceipt.objects.get();self.assertEqual(len(receipt.verification_code),6)
        self.assertTrue(receipt.verification_code.isascii() and receipt.verification_code.isalnum() and receipt.verification_code == receipt.verification_code.upper())
        response=self.client.get('/receipts/verify/',{'number':receipt.number,'code':receipt.verification_code})
        self.assertContains(response,'ANAK SATU')
        self.client.force_login(self.guardian)
        response=self.client.get(f'/receipts/{receipt.pk}/download.pdf')
        self.assertEqual(response.status_code,200);self.assertTrue(response.content.startswith(b'%PDF'))
        self.assertEqual(response['Cache-Control'],'private, no-store')
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(f'/receipts/{receipt.pk}/download.pdf').status_code,404)
        self.assertNotContains(self.client.get('/receipts/verify/',{'number':receipt.number,'code':receipt.verification_code}),'ANAK SATU')

    def test_greeting_uses_jakarta_time_boundaries_and_escapes_name(self):
        from datetime import datetime, timezone as dt_timezone
        from .views import dashboard_greeting
        for hour,expected in ((4,'Selamat malam'),(5,'Selamat pagi'),(10,'Selamat pagi'),(11,'Selamat siang'),(14,'Selamat siang'),(15,'Selamat sore'),(17,'Selamat sore'),(18,'Selamat malam')):
            utc_hour=(hour-7)%24
            self.assertEqual(dashboard_greeting(datetime(2026,10,27,utc_hour,tzinfo=dt_timezone.utc)),expected)
        self.guardian.first_name='<Admin>';self.guardian.save()
        self.client.force_login(self.guardian)
        response=self.client.get('/dashboard/')
        self.assertContains(response,'&lt;Admin&gt;')
        self.assertNotContains(response,'<Admin>')

    def test_school_announcement_workflow_and_qr_scope(self):
        from .models import SchoolAnnouncement
        admin=get_user_model().objects.create_user(username='announce-admin',password=TEST_PASSWORD_3)
        principal=get_user_model().objects.create_user(username='announce-principal',password=TEST_PASSWORD_3)
        Membership.objects.create(user=admin,school=self.school,role='admin')
        Membership.objects.create(user=principal,school=self.school,role='principal')
        self.client.force_login(admin)
        response=self.client.post('/announcements/new/',{'school':self.school.pk,'title':'Activity','body':'First paragraph\n\nSecond paragraph','event_date':'2026-11-20','pdf_url':'https://example.test/info.pdf'})
        self.assertEqual(response.status_code,302)
        item=SchoolAnnouncement.objects.get()
        self.assertEqual(self.client.post(f'/announcements/{item.pk}/approve/').status_code,404)
        self.assertEqual(self.client.post(f'/announcements/{item.pk}/submit/').status_code,302)
        self.assertEqual(self.client.get(f'/announcements/{item.pk}/edit/').status_code,400)
        self.client.force_login(principal)
        self.assertEqual(self.client.post(f'/announcements/{item.pk}/approve/').status_code,302)
        item.refresh_from_db();self.assertEqual(item.status,'approved');self.assertEqual(item.approved_by,principal)
        self.assertEqual(self.client.get(f'/announcements/verify/{item.approval_reference}/qr.png').status_code,200)
        self.assertContains(self.client.get(f'/announcements/verify/{item.approval_reference}/'),'not a certified digital signature')
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(f'/announcements/verify/{item.approval_reference}/').status_code,404)
        self.assertEqual(self.client.post(f'/announcements/{item.pk}/delete/').status_code,404)
    def test_announcement_rejects_non_https_pdf(self):
        admin=get_user_model().objects.create_user(username='announce-admin2',password=TEST_PASSWORD_3)
        Membership.objects.create(user=admin,school=self.school,role='admin')
        self.client.force_login(admin)
        response=self.client.post('/announcements/new/',{'school':self.school.pk,'title':'Activity','body':'Text','event_date':'2026-11-20','pdf_url':'http://example.test/a.pdf'})
        self.assertContains(response,'Use an HTTPS link.')

    def test_announcement_actions_require_and_accept_csrf_tokens(self):
        from django.test import Client, override_settings
        from .models import SchoolAnnouncement
        user=get_user_model().objects.create_user(username='csrf-announcement',password=TEST_PASSWORD_4)
        Membership.objects.create(user=user,school=self.school,role='admin')
        client=Client(enforce_csrf_checks=True);client.force_login(user)
        origin='https://pv-ac9ad268c614-51b1306d-b202-4d34-b04b-346824493b23.style.dev'
        with override_settings(CSRF_TRUSTED_ORIGINS=[origin]):
            page=client.get('/announcements/new/',secure=True)
            self.assertContains(page,'csrfmiddlewaretoken')
            token=client.cookies['csrftoken'].value
            fields={'school':self.school.pk,'title':'Example','body':'Example paragraph','event_date':'2026-11-20','pdf_url':''}
            self.assertEqual(client.post('/announcements/new/',fields,secure=True,HTTP_ORIGIN=origin).status_code,403)
            fields['csrfmiddlewaretoken']=token
            self.assertEqual(client.post('/announcements/new/',fields,secure=True,HTTP_ORIGIN=origin).status_code,302)
            announcement=SchoolAnnouncement.objects.get(title='Example')
            self.assertEqual(client.post(f'/announcements/{announcement.pk}/submit/',{'csrfmiddlewaretoken':token},secure=True,HTTP_ORIGIN=origin).status_code,302)

    def test_student_xlsx_template_preview_commit_and_duplicate(self):
        from io import BytesIO
        from django.core.files.uploadedfile import SimpleUploadedFile
        from openpyxl import load_workbook
        admin=get_user_model().objects.create_user(username='xlsx-students-admin',password=TEST_PASSWORD_3)
        Membership.objects.create(user=admin,school=self.school,role='admin')
        self.client.force_login(admin)
        template=self.client.get('/students/template.xlsx');self.assertEqual(template.status_code,200)
        book=load_workbook(BytesIO(template.content));book.active.append(['00088','New student'])
        stream=BytesIO();book.save(stream)
        response=self.client.post('/students/import/',{'school':self.school.pk,'file':SimpleUploadedFile('students.xlsx',stream.getvalue())})
        self.assertContains(response,'00088');self.assertFalse(Student.objects.filter(nis='00088').exists())
        token=self.client.session['student_import_preview']['token']
        self.assertEqual(self.client.post('/students/import/commit/',{'token':token}).status_code,302)
        self.assertTrue(Student.objects.filter(school=self.school,nis='00088').exists())
        self.assertEqual(self.client.post('/students/import/commit/',{'token':token}).status_code,400)
        response=self.client.post('/students/import/',{'school':self.school.pk,'file':SimpleUploadedFile('students.xlsx',stream.getvalue())})
        self.assertContains(response,'NIS sudah ada')
    def test_student_import_scope_and_numeric_nis(self):
        from io import BytesIO
        from django.core.files.uploadedfile import SimpleUploadedFile
        from openpyxl import load_workbook
        admin=get_user_model().objects.create_user(username='xlsx-isolated',password=TEST_PASSWORD_3)
        Membership.objects.create(user=admin,school=self.school,role='admin')
        self.client.force_login(admin)
        book=load_workbook(BytesIO(self.client.get('/students/template.xlsx').content));book.active.append([123,'Student'])
        stream=BytesIO();book.save(stream)
        upload=lambda:SimpleUploadedFile('students.xlsx',stream.getvalue())
        self.assertEqual(self.client.post('/students/import/',{'school':self.other_school.pk,'file':upload()}).status_code,404)
        response=self.client.post('/students/import/',{'school':self.school.pk,'file':upload()})
        self.assertContains(response,'NIS wajib berupa teks')
        self.assertNotIn('student_import_preview',self.client.session)
        self.client.force_login(self.guardian)
        self.assertEqual(self.client.get('/students/template.xlsx').status_code,404)
    def test_cross_school_membership_cannot_edit_or_withdraw_proof(self):
        # Guardianship (student.guardians) is school-agnostic, so a linked guardian who is
        # only a member of another school can own a proof on a SHB invoice. School scope
        # must still deny, and the control below proves the 404 came from school scope.
        user=get_user_model().objects.create_user(username='cross-school-guardian',password=TEST_PASSWORD_3)
        self.student.guardians.add(user)
        Membership.objects.create(user=user,school=self.other_school,role='guardian')
        proof=TransferProof.objects.create(invoice=self.invoice,guardian=user,reported_amount=100000,transfer_date=date(2026,10,27),image=b'x',image_type='png')
        self.client.force_login(user)
        self.assertEqual(self.client.get(f'/proofs/{proof.pk}/edit/').status_code,404)
        self.assertEqual(self.client.post(f'/proofs/{proof.pk}/edit/',{'amount':'999999','transfer_date':'2026-10-28'}).status_code,404)
        self.assertEqual(self.client.post(f'/proofs/{proof.pk}/withdraw/').status_code,404)
        proof.refresh_from_db()
        self.assertEqual((proof.reported_amount,proof.transfer_date,proof.status),(100000,date(2026,10,27),'pending'))
        Membership.objects.create(user=user,school=self.school,role='guardian')
        self.assertEqual(self.client.post(f'/proofs/{proof.pk}/edit/',{'amount':'999999','transfer_date':'2026-10-28'}).status_code,302)
        proof.refresh_from_db();self.assertEqual(proof.reported_amount,999999)

    def test_cross_school_enrollment_proof_is_not_editable_or_withdrawable(self):
        # An enrollment whose classroom belongs to another school is the one shape where the
        # student's school and the invoice's school disagree; it must stay out of scope.
        year=AcademicYear.objects.create(school=self.other_school,label='2026/27')
        other_class=Classroom.objects.create(year=year,name='9A')
        bad=Enrollment.objects.create(student=self.student,classroom=other_class)
        invoice=Invoice.objects.create(enrollment=bad,reference='BAD-PROOF',description='CROSS SCHOOL',amount=100000,due_date=date(2026,11,1))
        proof=TransferProof.objects.create(invoice=invoice,guardian=self.guardian,reported_amount=100000,transfer_date=date(2026,10,27),image=b'x',image_type='png')
        self.client.force_login(self.guardian)
        self.assertEqual(self.client.get(f'/proofs/{proof.pk}/edit/').status_code,404)
        self.assertEqual(self.client.post(f'/proofs/{proof.pk}/edit/',{'amount':'120000','transfer_date':'2026-10-28'}).status_code,404)
        self.assertEqual(self.client.post(f'/proofs/{proof.pk}/withdraw/').status_code,404)
        proof.refresh_from_db();self.assertEqual((proof.reported_amount,proof.status),(100000,'pending'))
        # Control: the same proof becomes editable once the enrollment is in-scope again.
        bad.classroom=self.classroom;bad.save()
        self.assertEqual(self.client.post(f'/proofs/{proof.pk}/edit/',{'amount':'120000','transfer_date':'2026-10-28'}).status_code,302)
        proof.refresh_from_db();self.assertEqual(proof.reported_amount,120000)

    def test_revoked_membership_blocks_edit_and_withdraw(self):
        # A live session is not the authorization boundary; the active membership is.
        proof=TransferProof.objects.create(invoice=self.invoice,guardian=self.guardian,reported_amount=100000,transfer_date=date(2026,10,27),image=b'x',image_type='png')
        Membership.objects.filter(user=self.guardian,school=self.school,role='guardian').update(active=False)
        self.client.force_login(self.guardian)
        self.assertEqual(self.client.get(f'/proofs/{proof.pk}/edit/').status_code,404)
        self.assertEqual(self.client.post(f'/proofs/{proof.pk}/edit/',{'amount':'120000','transfer_date':'2026-10-28'}).status_code,404)
        self.assertEqual(self.client.post(f'/proofs/{proof.pk}/withdraw/').status_code,404)
        proof.refresh_from_db();self.assertEqual((proof.reported_amount,proof.status),(100000,'pending'))
        Membership.objects.filter(user=self.guardian,school=self.school,role='guardian').update(active=True)
        self.assertEqual(self.client.post(f'/proofs/{proof.pk}/edit/',{'amount':'120000','transfer_date':'2026-10-28'}).status_code,302)
        proof.refresh_from_db();self.assertEqual(proof.reported_amount,120000)

    def test_same_school_staff_cannot_edit_or_withdraw_but_admin_can_verify(self):
        # The class teacher and the school admin can read or verify this proof, so the 404s
        # below pin the edit/withdraw right to the owning guardian rather than to any role.
        admin=get_user_model().objects.create_user(username='school-admin-proofs',password=TEST_PASSWORD_3)
        Membership.objects.create(user=admin,school=self.school,role='admin')
        proof=TransferProof.objects.create(invoice=self.invoice,guardian=self.guardian,reported_amount=100000,transfer_date=date(2026,10,27),image=b'x',image_type='png')
        for user in (self.teacher,admin):
            with self.subTest(user=user.username):
                self.client.force_login(user)
                self.assertEqual(self.client.get(f'/proofs/{proof.pk}/edit/').status_code,404)
                self.assertEqual(self.client.post(f'/proofs/{proof.pk}/edit/',{'amount':'120000','transfer_date':'2026-10-28'}).status_code,404)
                self.assertEqual(self.client.post(f'/proofs/{proof.pk}/withdraw/').status_code,404)
        proof.refresh_from_db();self.assertEqual((proof.reported_amount,proof.status),(100000,'pending'))
        self.client.force_login(admin)
        self.assertEqual(self.client.post(f'/proofs/{proof.pk}/verify/').status_code,302)
        proof.refresh_from_db();self.assertEqual(proof.status,'verified');self.assertEqual(Payment.objects.count(),1)

    def test_other_school_admin_sees_no_proof_and_cannot_act_on_it(self):
        proof=TransferProof.objects.create(invoice=self.invoice,guardian=self.guardian,reported_amount=137137,transfer_date=date(2026,10,27),image=b'x',image_type='png')
        self.client.force_login(self.other)
        listing=self.client.get('/proofs/')
        self.assertEqual(listing.status_code,200)
        self.assertNotContains(listing,'137137');self.assertNotContains(listing,'Lihat foto')
        self.assertEqual(self.client.get(f'/proofs/{proof.pk}/edit/').status_code,404)
        self.assertEqual(self.client.post(f'/proofs/{proof.pk}/edit/',{'amount':'120000','transfer_date':'2026-10-28'}).status_code,404)
        self.assertEqual(self.client.post(f'/proofs/{proof.pk}/withdraw/').status_code,404)
        self.assertEqual(self.client.post(f'/proofs/{proof.pk}/verify/').status_code,404)
        self.assertEqual(self.client.get(f'/proofs/{proof.pk}/image/').status_code,404)
        proof.refresh_from_db();self.assertEqual((proof.reported_amount,proof.status),(137137,'pending'))
        self.client.force_login(self.guardian)
        self.assertContains(self.client.get('/proofs/'),'137137')

class SeedDemoCommandTests(TestCase):
    def setUp(self):
        from django.db import connection
        if connection.vendor != 'sqlite':
            self.skipTest('seed_demo diizinkan hanya pada SQLite development.')
    def seed(self, **kwargs):
        from io import StringIO
        from django.core.management import call_command
        out=StringIO()
        call_command('seed_demo',stdout=out,**kwargs)
        return out.getvalue()
    def test_seed_creates_scoped_preview_data_and_repeats_safely(self):
        self.assertIn('objek baru dibuat',self.seed())
        self.assertTrue(School.objects.filter(name='Sekolah Harapan Bangsa (contoh)').exists())
        self.assertEqual(Student.objects.count(),3)
        self.assertEqual(Invoice.objects.count(),3)
        # Lima peran dari enam akun demo; superadmin tidak memakai membership.
        self.assertEqual(Membership.objects.count(),5)
        self.assertEqual(Enrollment.objects.count(),3)
        self.client.force_login(get_user_model().objects.get(username='ortu@shb.sch.id'))
        dashboard=self.client.get('/dashboard/')
        self.assertContains(dashboard,'Murid Contoh 1')
        self.assertNotContains(dashboard,'Murid Contoh 3')
        self.client.force_login(get_user_model().objects.get(username='admin@shb.sch.id'))
        self.assertContains(self.client.get('/dashboard/'),'Murid Contoh 3')
        self.assertIn('sudah lengkap',self.seed())
        counts=(School.objects.count(),Student.objects.count(),Invoice.objects.count(),Membership.objects.count(),Enrollment.objects.count())
        self.assertEqual(counts,(1,3,3,5,3))
    def test_seed_uses_given_password(self):
        self.seed(password=TEST_PASSWORD_1)
        self.assertTrue(self.client.login(username='ortu@shb.sch.id',password=TEST_PASSWORD_1))
    def test_seed_refuses_production(self):
        from io import StringIO
        from django.core.management import call_command
        from django.core.management.base import CommandError
        with self.settings(PRODUCTION=True):
            with self.assertRaises(CommandError):
                call_command('seed_demo',stdout=StringIO())
        self.assertFalse(School.objects.exists())

class ConcurrentProofVerificationTests(TransactionTestCase):
    """Verifikasi bukti bersamaan tidak boleh menghasilkan dua pembayaran.

    Memakai TransactionTestCase agar data benar-benar ter-commit dan terlihat oleh
    koneksi kedua. Kelas ini hanya berjalan pada PostgreSQL: SQLite mengunci seluruh
    database dan mengabaikan SELECT ... FOR UPDATE, sehingga tidak bisa menguji dua
    penulis bersamaan maupun perilaku menunggu lock.
    """
    def setUp(self):
        from django.db import connection
        if connection.vendor!='postgresql':
            self.skipTest('Butuh row lock PostgreSQL; SQLite mengabaikan SELECT ... FOR UPDATE.')
        User=get_user_model()
        self.school=School.objects.create(name='SHB Concurrent')
        year=AcademicYear.objects.create(school=self.school,label='2026/27')
        classroom=Classroom.objects.create(year=year,name='9A')
        guardian=User.objects.create_user(username='concurrent-guardian',password=TEST_PASSWORD_3)
        student=Student.objects.create(school=self.school,nis='900',name='ANAK KONKUREN')
        student.guardians.add(guardian)
        enrollment=Enrollment.objects.create(student=student,classroom=classroom)
        self.invoice=Invoice.objects.create(enrollment=enrollment,reference='INV-CONCURRENT',description='SPP',amount=200000,due_date=date(2026,11,1))
        self.admin=User.objects.create_user(username='concurrent-admin',password=TEST_PASSWORD_3)
        Membership.objects.create(user=self.admin,school=self.school,role='admin')
        self.proof=TransferProof.objects.create(invoice=self.invoice,guardian=guardian,reported_amount=100000,transfer_date=date(2026,10,27),image=b'x',image_type='png')
    def test_two_simultaneous_verifications_create_one_payment(self):
        import threading
        from django.db import connections
        from django.test import Client
        barrier=threading.Barrier(2)
        codes=[];errors=[]
        def verify():
            try:
                client=Client()
                client.force_login(self.admin)
                barrier.wait(timeout=15)
                codes.append(client.post(f'/proofs/{self.proof.pk}/verify/').status_code)
            except Exception as exc:
                errors.append(exc)
            finally:
                connections.close_all()
        threads=[threading.Thread(target=verify) for _ in range(2)]
        for thread in threads: thread.start()
        for thread in threads: thread.join(timeout=60)
        self.assertEqual(errors,[])
        self.assertEqual(sorted(codes),[302,400])
        self.assertEqual(Payment.objects.count(),1)
        self.proof.refresh_from_db()
        self.assertEqual(self.proof.status,TransferProof.Status.VERIFIED)
        self.assertEqual(self.proof.reviewed_by_id,self.admin.pk)
    def test_second_verification_waits_for_the_locked_row(self):
        """Verifikasi kedua benar-benar menunggu lock lalu melihat status ter-commit.

        Transaksi pertama memegang lock baris (sama seperti select_for_update pada view)
        dan menahannya sampai di-commit. Selama itu verifikasi kedua harus menunggu, bukan
        lolos; setelah commit ia melihat status VERIFIED dan menolak (400).
        """
        import threading
        import time
        from django.db import connections, transaction
        from django.test import Client
        locked=threading.Event();release=threading.Event();order=[];errors=[]
        def hold_row_lock_and_verify():
            try:
                with transaction.atomic():
                    proof=TransferProof.objects.select_for_update().get(pk=self.proof.pk)
                    locked.set()
                    release.wait(timeout=20)
                    Payment.objects.create(invoice=self.invoice,amount=proof.reported_amount,verified=True)
                    proof.status=TransferProof.Status.VERIFIED
                    proof.reviewed_by=self.admin
                    proof.save(update_fields=['status','reviewed_by'])
                    order.append('first-commit')
            except Exception as exc:
                errors.append(exc)
            finally:
                connections.close_all()
        def second_verification():
            try:
                client=Client()
                client.force_login(self.admin)
                order.append(('second',client.post(f'/proofs/{self.proof.pk}/verify/').status_code))
            except Exception as exc:
                errors.append(exc)
            finally:
                connections.close_all()
        holder=threading.Thread(target=hold_row_lock_and_verify)
        holder.start()
        self.assertTrue(locked.wait(timeout=20),'transaksi pertama gagal memegang lock baris')
        second=threading.Thread(target=second_verification)
        second.start()
        time.sleep(0.5)
        # Permintaan kedua harus tertahan selama transaksi pertama masih terbuka.
        # Asersi penentu ada di bawah: tanpa select_for_update, verifikasi kedua tetap
        # tertahan pada UPDATE namun tetap lolos (302) karena pemeriksaan status sudah
        # dilakukan sebelum menunggu, sehingga muncul dua pembayaran.
        self.assertEqual([item for item in order if isinstance(item,tuple)],[],
                         'permintaan verifikasi kedua sudah selesai padahal transaksi pertama masih terbuka')
        release.set()
        holder.join(timeout=30);second.join(timeout=30)
        self.assertEqual(errors,[])
        self.assertEqual(order[0],'first-commit')
        self.assertEqual(order[-1],('second',400))
        self.assertEqual(Payment.objects.count(),1)
        self.proof.refresh_from_db()
        self.assertEqual(self.proof.status,TransferProof.Status.VERIFIED)


class AsgiChunkedMultipartTests(TestCase):
    """The preview proxy forwards form bodies chunked and without Content-Length.

    Django's MultiPartParser takes the body size from CONTENT_LENGTH and returns an
    empty POST when the header is absent (django/http/multipartparser.py), so every
    multipart/form-data post — the profile form and every upload form — failed CSRF
    with "CSRF token missing" while urlencoded posts kept working. config/asgi.py
    forwards the real length; these tests pin that behaviour and prove CSRF is still
    enforced.
    """
    boundary='shb-test-boundary'
    origin='https://pv-ac9ad268c614-51b1306d-b202-4d34-b04b-346824493b23.style.dev'
    secret='a'*32

    def multipart_body(self, token=None):
        fields=[('first_name',b'Budi'),('last_name',b'Test')]
        if token is not None:
            fields.insert(0,('csrfmiddlewaretoken',token.encode()))
        body=b''
        for name,value in fields:
            body+=('--%s\r\nContent-Disposition: form-data; name="%s"\r\n\r\n'%(self.boundary,name)).encode()+value+b'\r\n'
        body+=('--%s\r\nContent-Disposition: form-data; name="photo"; filename=""\r\nContent-Type: application/octet-stream\r\n\r\n\r\n--%s--\r\n'%(self.boundary,self.boundary)).encode()
        return body

    def scope(self, headers, path='/accounts/profile/'):
        return {'type':'http','asgi':{'version':'3.0'},'http_version':'1.1','method':'POST','scheme':'http',
                'path':path,'raw_path':path.encode(),'query_string':b'','root_path':'',
                'server':('testserver',80),'client':('127.0.0.1',1234),'headers':headers}

    def multipart_headers(self, origin=None, content_type=None):
        return [(b'host',b'testserver'),
                (b'content-type',(content_type or 'multipart/form-data; boundary=%s'%self.boundary).encode()),
                (b'cookie',('csrftoken=%s'%self.secret).encode()),
                (b'origin',(origin or self.origin).encode()),
                (b'referer',((origin or self.origin)+'/accounts/profile/').encode()),
                (b'transfer-encoding',b'chunked')]

    def drive(self, application, scope, chunks):
        from asgiref.sync import async_to_sync
        from asgiref.testing import ApplicationCommunicator
        async def run():
            communicator=ApplicationCommunicator(application,scope)
            for body,more in chunks:
                await communicator.send_input({'type':'http.request','body':body,'more_body':more})
            message=await communicator.receive_output(timeout=5)
            while message['type']!='http.response.start':
                message=await communicator.receive_output(timeout=5)
            status=message['status']
            body=b''
            while True:
                message=await communicator.receive_output(timeout=5)
                body+=message.get('body',b'')
                if not message.get('more_body'):
                    break
            return status,body
        return async_to_sync(run)()

    def test_chunked_multipart_post_reaches_django_with_csrf(self):
        from config.asgi import application
        body=self.multipart_body(token=self.secret)
        half=len(body)//2
        status,response=self.drive(application,self.scope(self.multipart_headers()),
                                   [(body[:half],True),(body[half:],False)])
        # CSRF passed, so the pipeline now redirects the anonymous post to login.
        self.assertEqual(status,302)

    def test_chunked_multipart_post_without_token_is_still_rejected(self):
        from config.asgi import application
        body=self.multipart_body()
        half=len(body)//2
        status,response=self.drive(application,self.scope(self.multipart_headers()),
                                   [(body[:half],True),(body[half:],False)])
        self.assertEqual(status,403)
        self.assertIn(b'CSRF verification failed',response)

    def test_chunked_multipart_post_from_untrusted_origin_is_rejected(self):
        from config.asgi import application
        body=self.multipart_body(token=self.secret)
        half=len(body)//2
        status,response=self.drive(application,self.scope(self.multipart_headers(origin='https://evil.example.test')),
                                   [(body[:half],True),(body[half:],False)])
        self.assertEqual(status,403)
        self.assertIn(b'CSRF verification failed',response)

    def test_requests_without_the_problem_are_passed_through_untouched(self):
        from config.asgi import ChunkedMultipartBody
        seen={}
        async def inner(scope,receive,send):
            seen['scope']=scope
            seen['messages']=[]
            while True:
                message=await receive()
                seen['messages'].append(message)
                if message.get('type')!='http.request' or not message.get('more_body'):
                    break
            await send({'type':'http.response.start','status':200,'headers':[]})
            await send({'type':'http.response.body','body':b''})
        bridge=ChunkedMultipartBody(inner)
        cases={'already has a length':[(b'content-length',b'12'),(b'content-type',b'multipart/form-data; boundary=x')],
               'chunked but urlencoded':[(b'content-type',b'application/x-www-form-urlencoded'),(b'transfer-encoding',b'chunked')],
               'chunked multipart on GET':[(b'content-type',b'multipart/form-data; boundary=x'),(b'transfer-encoding',b'chunked')]}
        for label,headers in cases.items():
            scope=self.scope(list(headers),path='/accounts/profile/')
            if label.endswith('GET'): scope['method']='GET'
            status,_=self.drive(bridge,scope,[(b'first_name=Budi&last_name=Test',False)])
            self.assertEqual(status,200,'the inner app must answer untouched: %s'%label)
            self.assertEqual([value for name,value in seen['scope']['headers'] if name==b'content-length'],
                             [value for name,value in scope['headers'] if name==b'content-length'],'unexpected length injected for: %s'%label)
            self.assertEqual(seen['messages'][0]['body'],b'first_name=Budi&last_name=Test','body must reach the app: %s'%label)

    def test_chunked_multipart_body_is_replayed_whole_in_blocks(self):
        from unittest.mock import patch
        from config.asgi import ChunkedMultipartBody
        seen={}
        async def inner(scope,receive,send):
            seen['scope']=scope
            body=b''
            while True:
                message=await receive()
                body+=message.get('body',b'')
                if not message.get('more_body'):
                    break
            seen['body']=body
            await send({'type':'http.response.start','status':200,'headers':[]})
            await send({'type':'http.response.body','body':b''})
        body=self.multipart_body(token=self.secret)
        third=len(body)//3
        chunks=[(body[:third],True),(body[third:2*third],True),(body[2*third:],False)]
        with patch('config.asgi.READ_BLOCK',8):
            self.drive(ChunkedMultipartBody(inner),self.scope(self.multipart_headers()),chunks)
        self.assertEqual(seen['body'],body,'the full body must arrive unchanged')
        self.assertIn((b'content-length',str(len(body)).encode()),seen['scope']['headers'],'the real length must be published')


